"""Dump reference output workbook analysis block to confirm exact labels/order."""
import openpyxl, sys
wb = openpyxl.load_workbook(r'C:\Users\TEST\Documents\Reference\Output ExcelSpreadsheet.xlsx', data_only=True)
ws = wb['Table 1']
print(f"max_row={ws.max_row}, max_col={ws.max_column}")
print("\nFirst 10 rows:")
for r in range(1, 11):
    vals = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column+1)]
    print(f"  Row {r}: {vals}")
print("\nLast 30 rows:")
for r in range(max(1, ws.max_row-29), ws.max_row+1):
    vals = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column+1)]
    print(f"  Row {r}: {vals}")

# Also dump formulas (non-data view)
wb2 = openpyxl.load_workbook(r'C:\Users\TEST\Documents\Reference\Output ExcelSpreadsheet.xlsx', data_only=False)
ws2 = wb2['Table 1']
print("\n\nFormulas in last 30 rows:")
for r in range(max(1, ws2.max_row-29), ws2.max_row+1):
    vals = [ws2.cell(row=r, column=c).value for c in range(1, ws2.max_column+1)]
    has_content = any(v is not None for v in vals)
    if has_content:
        print(f"  Row {r}: {vals}")

# Also dump source workbook analysis block
print("\n\n=== SOURCE WORKBOOK analysis rows ===")
wb3 = openpyxl.load_workbook(r'C:\Users\TEST\Documents\Reference\Original ExcelSpreadsheet.xlsx', data_only=False)
ws3 = wb3['Table 1']
print(f"max_row={ws3.max_row}, max_col={ws3.max_column}")
for r in range(49, ws3.max_row+1):
    vals = [ws3.cell(row=r, column=c).value for c in range(1, ws3.max_column+1)]
    has_content = any(v is not None for v in vals)
    if has_content:
        print(f"  Row {r}: {vals}")
