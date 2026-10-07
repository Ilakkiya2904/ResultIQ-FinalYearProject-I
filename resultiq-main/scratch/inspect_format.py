import openpyxl
from openpyxl.utils import get_column_letter

def rgb_of(color):
    try:
        if color and color.type == 'rgb':
            return color.rgb
    except Exception:
        pass
    return '?'

def inspect_formatting(path, label):
    wb = openpyxl.load_workbook(path, data_only=False)
    ws = wb.active
    with open(r'c:\Users\TEST\Documents\ResultIQ-new\scratch\ref_format.txt', 'w', encoding='utf-8') as f:
        f.write(f'=== {label} ===\n')
        f.write(f'Sheet: {ws.title}  max_row={ws.max_row}  max_col={ws.max_column}\n')
        f.write(f'Merged cells: {[str(m) for m in ws.merged_cells.ranges]}\n\n')
        # Check row 49 onwards (analysis area)
        f.write('--- Analysis area (rows 49-60) ---\n')
        for r in range(49, 61):
            for c in range(1, ws.max_column+1):
                cell = ws.cell(r, c)
                if cell.value is not None or cell.has_style:
                    f.write(f'  Cell {get_column_letter(c)}{r}: val={repr(cell.value)}\n')
                    f.write(f'    font: bold={cell.font.bold}, size={cell.font.size}, color={rgb_of(cell.font.color)}\n')
                    f.write(f'    fill: type={cell.fill.fill_type}, fg={rgb_of(cell.fill.fgColor)}\n')
                    f.write(f'    border: left={cell.border.left.border_style}, right={cell.border.right.border_style}, top={cell.border.top.border_style}, bottom={cell.border.bottom.border_style}\n')
                    f.write(f'    align: h={cell.alignment.horizontal}, v={cell.alignment.vertical}, wrap={cell.alignment.wrap_text}\n')
                    f.write(f'    num_fmt={cell.number_format}\n')

        f.write('\n--- Column widths ---\n')
        for col_dim in ws.column_dimensions.values():
            f.write(f'  col {col_dim.index}: width={col_dim.width}\n')

        f.write('\n--- Row heights (rows 1-10 + 49-60) ---\n')
        for r in list(range(1, 11)) + list(range(49, 61)):
            if r in ws.row_dimensions:
                f.write(f'  row {r}: height={ws.row_dimensions[r].height}\n')

        f.write('\n--- Header row (row 3) cells ---\n')
        for c in range(1, ws.max_column+1):
            cell = ws.cell(3, c)
            f.write(f'  {get_column_letter(c)}3: val={repr(cell.value)}, bold={cell.font.bold}, fill_fg={rgb_of(cell.fill.fgColor)}, border_l={cell.border.left.border_style}\n')

    wb.close()

inspect_formatting(
    r'C:\Users\TEST\Documents\ResultIQ-new\Reference\Output ExcelSpreadsheet.xlsx',
    'OUTPUT FORMAT'
)
print('done')
